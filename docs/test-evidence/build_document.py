from docx import Document
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path
root=Path(__file__).resolve().parent
out=root.parent/'Test Evidence.docx'
doc=Document()
s=doc.sections[0]
s.top_margin=s.bottom_margin=Inches(.65)
s.left_margin=s.right_margin=Inches(.65)
for name in ['Normal','Title','Heading 1','Heading 2','Caption','Subtitle']:
 style=doc.styles[name]
 style.font.name='Calibri'
 style.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
from docx.oxml.ns import qn
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
 border.getparent().remove(border)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.add_paragraph('Demo 2 test evidence', 'Title')
doc.add_paragraph('DataBridge frontend and backend split', 'Subtitle')
doc.add_paragraph('23 October 2026')
doc.add_paragraph('All three automated tests passed against the completed local demo. This report records the API and frontend results and syntax checks for the course presenter. The images are browser screenshots of the saved, unedited command output; the original text logs and HTML views are retained alongside the PNG files in test-evidence.')
doc.add_heading('Automated tests',1)
doc.add_picture(str(root/'test-run.png'),width=Inches(7.0))
doc.add_paragraph('Figure 1  API and frontend test run passed with three tests and zero failures. Exit code 0.', 'Caption')
doc.add_paragraph('The API test makes real local HTTP requests and checks the 12-product JSON catalog, combined search and category filters, empty results, CORS, malformed queries returning 400, unknown routes returning 404, and invalid storage data returning a generic 500 response. The three error messages in the output are expected fault-injection results. The original catalog is restored after these checks.')
doc.add_paragraph('The first frontend test checks literal text rendering, filtering, the empty state, and retry without duplicate categories. The second checks recovery from network errors, HTTP errors, and malformed JSON response shapes. These tests use a DOM test double and do not measure browser layout.')
doc.add_page_break()
doc.add_heading('Syntax and whitespace checks',1)
doc.add_picture(str(root/'syntax-check.png'),width=Inches(7.0))
doc.add_paragraph('Figure 2  Both JavaScript syntax checks and Git whitespace validation passed with exit code 0. The line-ending warning is informational.', 'Caption')
doc.add_heading('Manual checks still required',1)
doc.add_paragraph('Responsive browser layout at 320 pixels and desktop widths, screen-reader announcements, and the browser Network-panel walkthrough have not been visually verified in this run. The README provides the manual steps. No screenshot in this report claims that those checks passed.')
doc.add_heading('Reproduce the checks',1)
doc.add_paragraph('From the demo directory, run npm ci --prefix backend and npm test --prefix backend. Run node --check frontend/script.js, node --check backend/routes/items.js, and git diff --check for the additional checks. Use Node.js 22 or newer. Run tests separately from a live classroom server because the API test briefly changes and restores the catalog file.')
doc.add_paragraph('Implementation commit: ef73d3c. Evidence is recorded in a subsequent local commit. No remote push or deployment was performed.')
doc.save(out)
print(out)

